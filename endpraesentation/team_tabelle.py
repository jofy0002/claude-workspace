"""Erzeugt Team_Leitfaden_Elbharz.xlsx – schlichte gemeinsame Abhak-Liste fürs Team (Excel/Teams, ohne Claude nutzbar)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from datetime import time

NAMES = ["Jona", "Lennart", "Sönke", "Nio"]
PETROL, AMBER, INK, MUTED = "0F5C6B", "D9822B", "1B2A2F", "7C8C90"
FIELD = PatternFill("solid", fgColor="F1F0EB")      # Feld zum Anklicken
DONE = PatternFill("solid", fgColor="DDEFC6")       # erledigt
LINE = Side(style="thin", color="E2E0D8")
F = lambda **k: Font(name="Arial", **{"color": INK, "size": 12, **k})
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

def sheet_base(ws, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "portrait"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

def heading(ws, row, text, sub):
    ws.cell(row, 2, text).font = F(size=14, bold=True, color=PETROL)
    ws.cell(row + 1, 2, sub).font = F(size=11, color=MUTED)
    ws.row_dimensions[row].height = 26

def colhead(ws, row, cols):
    for col, text in cols:
        c = ws.cell(row, col, text); c.font = F(size=11, bold=True, color=MUTED)
        c.alignment = CENTER if col > 2 else LEFT
        c.border = Border(bottom=Side(style="medium", color=PETROL))

wb = Workbook()

# ================================================================ Aufgaben
ws = wb.active; ws.title = "Aufgaben"
sheet_base(ws, {"A": 2, "B": 64, "C": 13, "D": 13, "E": 13, "F": 13, "G": 2})
ws["B1"] = "Elbharz · Aufgaben bis zur Präsentation"; ws["B1"].font = F(size=20, bold=True, color=PETROL)
ws.row_dimensions[1].height = 34
ws["B2"] = "Graues Feld anklicken → Pfeil → ✓ oder deinen Namen wählen. Erledigtes wird grün."
ws["B2"].font = F(size=11, color=MUTED)

einzeln = [
    "Sprechtext für meine Folien in eigene Worte gebracht",
    "Meinen Teil laut mit Stoppuhr geübt",
    "Kann die Zahlen meiner Folien ohne Ablesen erklären",
    "Reflexion (Folie 10) gelesen – passt für mich",
    "Generalprobe mitgemacht",
]
gemeinsam = [
    "Karte Norddeutschland mit roter Elbe einfügen (Folie 3)",
    "KI-Teambild einfügen + Tool notieren (Folie 12)",
    "Bildquelle Limette & Tamarinde klären (Quellenfolie)",
    "R3 klären: 18,2 Paletten und Mitarbeiterzahl laut Protokoll",
    "Präsentation am Schul-Beamer testen",
    "Alle Platzhalter ersetzt („BILD FOLGT“, „Quelle ergänzen“)",
    "PDF exportieren + Backup auf USB-Stick",
    "Präsentation, PDF und Bögen M15 an die Lehrkraft",
]

# --- Teil 1: jede:r für sich
r0 = 5
heading(ws, r0, "Jede:r für sich", "In deiner eigenen Spalte ✓ wählen.")
colhead(ws, r0 + 2, [(2, "Aufgabe")] + [(3 + i, n) for i, n in enumerate(NAMES)])
dv_check = DataValidation(type="list", formula1='"✓"', allow_blank=True)
dv_check.showErrorMessage = True; dv_check.error = "Bitte ✓ auswählen oder Feld leeren."
ws.add_data_validation(dv_check)
a1 = r0 + 3
for k, t in enumerate(einzeln):
    r = a1 + k
    ws.row_dimensions[r].height = 30
    c = ws.cell(r, 2, t); c.font = F(); c.alignment = LEFT; c.border = Border(bottom=LINE)
    for col in range(3, 7):
        f = ws.cell(r, col); f.fill = FIELD; f.alignment = CENTER; f.font = F(size=14, bold=True, color="3D6610")
        f.border = Border(bottom=Side(style="thin", color="FFFFFF"), left=Side(style="thin", color="FFFFFF"), right=Side(style="thin", color="FFFFFF"))
        dv_check.add(f.coordinate)
a2 = a1 + len(einzeln) - 1
ws.conditional_formatting.add(f"C{a1}:F{a2}", CellIsRule(operator="equal", formula=['"✓"'], fill=DONE))
ws.conditional_formatting.add(f"B{a1}:B{a2}", FormulaRule(formula=[f'COUNTIF($C{a1}:$F{a1},"✓")=4'], fill=DONE))

# --- Teil 2: eine:r reicht
r1 = a2 + 3
heading(ws, r1, "Eine:r reicht", "Wer es erledigt hat, wählt seinen Namen.")
colhead(ws, r1 + 2, [(2, "Aufgabe"), (3, "Erledigt von")])
ws.merge_cells(start_row=r1 + 2, start_column=3, end_row=r1 + 2, end_column=4)
dv_name = DataValidation(type="list", formula1='"' + ",".join(NAMES) + '"', allow_blank=True)
ws.add_data_validation(dv_name)
b1 = r1 + 3
for k, t in enumerate(gemeinsam):
    r = b1 + k
    ws.row_dimensions[r].height = 30
    c = ws.cell(r, 2, t); c.font = F(); c.alignment = LEFT; c.border = Border(bottom=LINE)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    f = ws.cell(r, 3); f.fill = FIELD; f.alignment = CENTER; f.font = F(bold=True, color=PETROL)
    white = Side(style="thin", color="FFFFFF")
    for col in (3, 4): ws.cell(r, col).border = Border(top=white, bottom=white)
    dv_name.add(f.coordinate)
b2 = b1 + len(gemeinsam) - 1
ws.conditional_formatting.add(f"B{b1}:D{b2}", FormulaRule(formula=[f'$C{b1}<>""'], fill=DONE))

# --- Fortschritt oben
total = len(einzeln) * 4 + len(gemeinsam)
ws["B3"] = f'="Fortschritt: "&(COUNTIF(C{a1}:F{a2},"✓")+COUNTA(C{b1}:C{b2}))&" von {total} erledigt"'
ws["B3"].font = F(size=12, bold=True, color=AMBER)
ws.row_dimensions[3].height = 22
ws.freeze_panes = "A4"

# ================================================================ Wer spricht was
zp = wb.create_sheet("Wer spricht was")
sheet_base(zp, {"A": 2, "B": 8, "C": 42, "D": 16, "E": 10, "F": 2})
zp["B1"] = "Wer spricht was"; zp["B1"].font = F(size=20, bold=True, color=PETROL)
zp.row_dimensions[1].height = 34
zp["B2"] = "Pro Folie einen Namen wählen – Wechsel sind jederzeit möglich."; zp["B2"].font = F(size=11, color=MUTED)
colhead(zp, 4, [(2, "Folie"), (3, "Titel"), (4, "Wer?"), (5, "Dauer")])
zp.cell(4, 2).alignment = CENTER; zp.cell(4, 3).alignment = LEFT
plan = [("Titel", 0, 20), ("Wer wir sind", 0, 45), ("Nachhaltigkeit", 0, 45), ("Produziert, angeboten, verkauft", 1, 0),
        ("Warum wir so entschieden haben", 1, 5), ("Wirtschaftsdienst", 0, 50), ("Unsere Kennzahlen", 1, 0),
        ("Wirtschaftlichkeit & Schlussbestand", 0, 55), ("Verkaufswert", 0, 50), ("Reflexion", 1, 20),
        ("Bezug zur Gründung", 1, 10), ("Vielen Dank", 0, 10)]
dv_who = DataValidation(type="list", formula1='"' + ",".join(NAMES) + '"', allow_blank=True)
zp.add_data_validation(dv_who)
for k, (t, m, s) in enumerate(plan):
    r = 5 + k
    zp.row_dimensions[r].height = 26
    a = zp.cell(r, 2, k + 1); a.font = F(color=MUTED); a.alignment = CENTER
    b = zp.cell(r, 3, t); b.font = F(); b.alignment = LEFT
    w = zp.cell(r, 4); w.fill = FIELD; w.alignment = CENTER; w.font = F(bold=True, color=PETROL); dv_who.add(w.coordinate)
    d = zp.cell(r, 5, time(0, m, s)); d.number_format = "m:ss"; d.font = F(color=MUTED); d.alignment = CENTER
    for col in (2, 3, 5): zp.cell(r, col).border = Border(bottom=LINE)
zp.conditional_formatting.add("B5:E16", FormulaRule(formula=['$D5<>""'], fill=DONE))

zp.cell(18, 3, "Zeit pro Person").font = F(size=14, bold=True, color=PETROL)
for i, n in enumerate(NAMES):
    r = 19 + i
    zp.cell(r, 3, n).font = F(); zp.cell(r, 3).alignment = LEFT
    c = zp.cell(r, 4, f'=SUMIF($D$5:$D$16,"{n}",$E$5:$E$16)'); c.number_format = "m:ss"; c.font = F(bold=True); c.alignment = CENTER
zp.cell(23, 3, "Gesamt (Ziel 10–12 Min.)").font = F(bold=True)
c = zp.cell(23, 4, "=SUM(E5:E16)"); c.number_format = "m:ss"; c.font = F(bold=True, color=AMBER); c.alignment = CENTER
zp.cell(24, 3, "Noch ohne Sprecher:in").font = F(color=MUTED)
c = zp.cell(24, 4, '=COUNTBLANK(D5:D16)&" Folien"'); c.font = F(color=MUTED); c.alignment = CENTER

ws.sheet_properties.tabColor = PETROL; zp.sheet_properties.tabColor = AMBER
wb.active = 0
wb.save("Team_Leitfaden_Elbharz.xlsx")
print("ok")
