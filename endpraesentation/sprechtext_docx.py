"""Erzeugt Sprechtext_Elbharz.docx aus den Notizen der Präsentation."""
from pptx import Presentation
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

PETROL, AMBER, MUTED = RGBColor(0x0F, 0x5C, 0x6B), RGBColor(0xD9, 0x82, 0x2B), RGBColor(0x5F, 0x8F, 0x99)
TITEL = ["Titel + Einstiegsfrage", "Wer wir sind", "Nachhaltigkeit", "Produziert, angeboten, verkauft", "Drei Entscheidungen",
         "Zwei Kennzahlen", "Wirtschaftlichkeit & Schlussbestand", "Verkaufswert", "Wirtschaftsdienst", "Reflexion",
         "Bezug zur Gründung", "Schluss + Danke"]
UEBERGABE = {3: "Übergabe an Folie 4: „… und wie das im Spiel lief, zeigt euch jetzt [Name].“",
             6: "Übergabe an Folie 7: „Was das in Zahlen bedeutet, erklärt [Name].“",
             9: "Übergabe an Folie 10: „Was wir daraus mitnehmen, sagt euch [Name].“"}

prs = Presentation("Elbharz_Endpraesentation.pptx")
doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Cm(2.2); sec.top_margin = sec.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(12)

h = doc.add_paragraph(); r = h.add_run("Sprechtext Endpräsentation"); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = PETROL
p = doc.add_paragraph(); r = p.add_run("Elbharz Getränke GbR · Klasse 11b · ca. 10:30 Minuten"); r.font.color.rgb = MUTED
tips = doc.add_paragraph()
r = tips.add_run("So nutzt ihr den Text: "); r.bold = True
tips.add_run("Das ist ein Vorschlag – sagt es in euren eigenen Worten. Genaue Zahlen stehen auf den Folien, im Sprechen reichen runde Zahlen. "
             "Nicht ablesen, was alle sehen, sondern erzählen, was ihr gedacht habt. Text in [eckigen Klammern] noch anpassen. "
             "Sprecher:innen tragt ihr in der Excel-Liste „Wer spricht was“ ein.")

for i, s in enumerate(list(prs.slides)[:12]):
    n = s.notes_slide.notes_text_frame.text
    zeit = n.splitlines()[0].replace("ZEIT: ", "").replace("ca. ", "").replace(" min", "")
    text = n.split("anpassen):\n")[-1].strip()
    hd = doc.add_paragraph(); hd.paragraph_format.space_before = Pt(16); hd.paragraph_format.keep_with_next = True
    r = hd.add_run(f"Folie {i+1} – {TITEL[i]}"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = PETROL
    meta = doc.add_paragraph(); meta.paragraph_format.keep_with_next = True
    r = meta.add_run(f"{zeit} Min.   ·   Sprecher:in: ______________"); r.font.size = Pt(11); r.font.color.rgb = MUTED
    doc.add_paragraph(text)
    if i + 1 in UEBERGABE:
        u = doc.add_paragraph(); r = u.add_run(UEBERGABE[i + 1]); r.italic = True; r.font.color.rgb = AMBER

doc.add_page_break()
h = doc.add_paragraph(); r = h.add_run("Falls die Lehrkraft nachfragt"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = PETROL
for q, a in [
    ("Wie habt ihr die Wirtschaftlichkeit berechnet?", "Umsatzerlöse geteilt durch alle Kosten der Runde, z. B. R3: 79.940 € ÷ 54.000 € = 1,48. Alle Werte: Anhang-Folie."),
    ("Wie aussagekräftig ist die Kennzahl?", "Maschinenkäufe zählen voll in der Kaufrunde. Rechnet man stattdessen in allen Runden nur den Wertverlust (10 % je Nutzungsrunde), ergibt sich: R1 1,07 · R2 1,16 · R3 1,33 · R4 0,22 · R5 0,09 · R6 ≈ 1,00."),
    ("Warum habt ihr in R2 eine Maschine ohne Mitarbeiter gekauft?", "Wir wollten Kapazität für R3 aufbauen; der Mitarbeiter kam erst in R3. Im Nachhinein hätten wir beides zusammen planen sollen."),
    ("Warum in R4/R5 so ein hoher Preis?", "Wir erwarteten weniger Absatz und wollten mit einer Palette kein Minus machen. Das war ein Denkfehler: Der Preis bestimmt den Absatz."),
    ("Wie kommt ihr auf 18.000 € Restwert für Maschine 2?", "10 % Wertverlust je genutzter Runde; sie lief nur R3–R6, also 4 Runden: 30.000 − 4 × 3.000."),
    ("Wie würde eure Spielerweiterung funktionieren?", "Marketingbudget: je 2.000 € → +1 Palette Absatz. Mehrweg-Option: +1.000 € Kosten je Runde → +1 Palette Absatz."),
]:
    p = doc.add_paragraph(); r = p.add_run(q); r.bold = True
    doc.add_paragraph(a)
doc.save("Sprechtext_Elbharz.docx"); print("ok")
